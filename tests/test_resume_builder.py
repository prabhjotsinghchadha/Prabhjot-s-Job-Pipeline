"""Tests for utils/resume_builder.py, utils/ats_checker.py and the structured
tailoring validation in utils/resume_tailor.py — no LLM, no network."""
import io

import pytest

from utils.resume_builder import (
    apply_tailoring,
    build_resume,
    has_structured_resume,
    render_docx,
    render_pdf,
    render_text,
    resume_from_profile,
    resume_path_for_job,
    safe_filename,
)
from utils.ats_checker import (
    check_ats,
    check_ats_file,
    extract_keywords,
    extract_pdf_text,
    keywords_present,
)
from utils.resume_tailor import _validate_structured, _with_defaults


PROFILE = {
    "personal": {
        "first_name": "Ada", "last_name": "Lovelace", "email": "ada@example.com",
        "phone": "+1 555 000 1234", "location": "Remote", "linkedin": "https://linkedin.com/in/ada",
        "github": "https://github.com/ada", "portfolio": "https://ada.dev/",
    },
    "resume_path": "",
    "skills": {"primary": ["Python"], "secondary": ["AWS"]},
    "preferences": {"keywords": ["TypeScript", "Next.js"]},
    "resume": {
        "headline": "Senior Full Stack Engineer",
        "summary": "Engineer with 8+ years shipping TypeScript and Node.js products on GCP.",
        "skills": [
            {"category": "Languages", "items": ["TypeScript", "Python"]},
            {"category": "Cloud", "items": ["GCP", "AWS", "Docker"]},
        ],
        "experience": [
            {"company": "Acme", "title": "Senior Engineer", "location": "Remote",
             "start": "Jan 2022", "end": "Present", "tagline": "b2b saas",
             "bullets": ["Built a Next.js platform serving 20,000+ users.",
                         "Reduced p95 latency by 40% with Redis caching."]},
            {"company": "Globex", "title": "Engineer", "location": "Berlin",
             "start": "2018", "end": "2021", "bullets": "Shipped REST APIs on AWS.\nLed a team of 3."},
        ],
        "education": [{"degree": "B.S. Computer Science", "school": "State U",
                       "location": "City", "start": "2014", "end": "2018"}],
        "projects": [{"name": "Tool A", "url": "https://a.dev", "tech": ["Rust"]},
                     {"name": "Tool B", "description": "cli"}],
        "certifications": ["AWS SAA", {"name": "GCP ACE", "issuer": "Google", "year": "2023"}],
        "languages": ["English"],
    },
}

POSTING = """
Senior Full Stack Engineer (Remote)
We're looking for a senior engineer to own our Next.js + Node.js product.
Requirements:
- 5+ years with TypeScript, React and Node.js
- Experience with GCP or AWS, Docker, CI/CD
- Real-time systems (WebSockets) and PostgreSQL
- Bonus: LLM integration, Kubernetes
Benefits: equity, remote-first.
"""


# ─── normalisation ───────────────────────────────────────────────────────────

class TestResumeFromProfile:
    def test_contact_merged_from_personal(self):
        r = resume_from_profile(PROFILE)
        assert r["contact"]["name"] == "Ada Lovelace"
        assert r["contact"]["email"] == "ada@example.com"

    def test_bullets_accept_newline_string(self):
        r = resume_from_profile(PROFILE)
        assert r["experience"][1]["bullets"] == ["Shipped REST APIs on AWS.", "Led a team of 3."]
        assert r["experience"][1]["end"] == "2021"

    def test_missing_end_defaults_to_present(self):
        p = {"personal": {}, "resume": {"experience": [{"company": "X", "title": "Y", "start": "2020"}]}}
        assert resume_from_profile(p)["experience"][0]["end"] == "Present"

    def test_certifications_accept_strings_and_dicts(self):
        r = resume_from_profile(PROFILE)
        assert r["certifications"][0]["name"] == "AWS SAA"
        assert r["certifications"][1]["issuer"] == "Google"

    def test_skills_fallback_to_profile_skills(self):
        p = {"personal": {}, "skills": {"primary": ["Go"], "secondary": ["K8s"]}, "resume": {"summary": "x"}}
        r = resume_from_profile(p)
        assert [g["items"] for g in r["skills"]] == [["Go"], ["K8s"]]

    def test_has_structured_resume(self):
        assert has_structured_resume(PROFILE)
        assert not has_structured_resume({"resume_path": "x.pdf"})
        assert not has_structured_resume({"resume": {"projects": []}})


# ─── tailoring overlay ───────────────────────────────────────────────────────

class TestApplyTailoring:
    def test_none_is_noop(self):
        base = resume_from_profile(PROFILE)
        assert apply_tailoring(base, None) == base

    def test_overrides_summary_headline_and_bullets(self):
        base = resume_from_profile(PROFILE)
        out = apply_tailoring(base, {
            "tailored_headline": "Staff Engineer",
            "tailored_summary": "New summary.",
            "tailored_experience": [{"index": 0, "bullets": ["Rewritten."]}],
        })
        assert out["headline"] == "Staff Engineer"
        assert out["summary"] == "New summary."
        assert out["experience"][0]["bullets"] == ["Rewritten."]
        assert out["experience"][1]["bullets"] == base["experience"][1]["bullets"]
        # base untouched
        assert base["experience"][0]["bullets"][0].startswith("Built")

    def test_ignores_bad_indexes_and_garbage(self):
        base = resume_from_profile(PROFILE)
        out = apply_tailoring(base, {
            "tailored_experience": [{"index": 9, "bullets": ["x"]}, "junk", {"index": "a"}],
            "tailored_skills": ["not a dict", {"category": "Empty", "items": []}],
            "project_order": [1, 1, 7, "x"],
        })
        assert out["experience"] == base["experience"]
        assert out["skills"] == base["skills"]
        assert [p["name"] for p in out["projects"]] == ["Tool B", "Tool A"]

    def test_never_changes_company_or_dates(self):
        base = resume_from_profile(PROFILE)
        out = apply_tailoring(base, {"tailored_experience": [{"index": 0, "bullets": ["a"], "company": "Evil"}]})
        assert out["experience"][0]["company"] == "Acme"
        assert out["experience"][0]["start"] == "Jan 2022"


class TestValidateStructured:
    def test_caps_bullets_at_original_plus_one(self):
        base = resume_from_profile(PROFILE)
        result = _with_defaults({"tailored_experience": [{"index": 0, "bullets": ["1", "2", "3", "4", "5"]}]})
        out = _validate_structured(result, base)
        assert len(out["tailored_experience"][0]["bullets"]) == 3
        assert out["tailored_bullets"] == ["1", "2", "3"]

    def test_drops_out_of_range(self):
        base = resume_from_profile(PROFILE)
        out = _validate_structured(_with_defaults({"tailored_experience": [{"index": 5, "bullets": ["x"]}]}), base)
        assert out["tailored_experience"] == []


# ─── renderers ───────────────────────────────────────────────────────────────

class TestRenderers:
    def test_text_has_sections_in_order(self):
        text = render_text(resume_from_profile(PROFILE))
        order = ["ADA LOVELACE", "PROFESSIONAL SUMMARY", "CORE SKILLS", "PROFESSIONAL EXPERIENCE",
                 "PROJECTS", "EDUCATION", "CERTIFICATIONS", "LANGUAGES"]
        positions = [text.index(s) for s in order]
        assert positions == sorted(positions)
        assert "• Built a Next.js platform" in text
        assert "linkedin.com/in/ada" in text and "https://" not in text.split("\n")[3]

    def test_text_skips_empty_sections(self):
        text = render_text(resume_from_profile(PROFILE))
        assert "AWARDS" not in text

    def test_docx_is_valid_and_contains_content(self):
        data = render_docx(resume_from_profile(PROFILE))
        assert data[:2] == b"PK"
        from docx import Document
        doc = Document(io.BytesIO(data))
        joined = "\n".join(p.text for p in doc.paragraphs)
        assert "PROFESSIONAL EXPERIENCE" in joined
        assert "Reduced p95 latency by 40%" in joined
        assert not doc.tables, "ATS rule: no tables in DOCX"
        assert doc.core_properties.author == "Ada Lovelace"

    def test_pdf_has_clean_text_layer(self, tmp_path):
        data = render_pdf(resume_from_profile(PROFILE))
        assert data.startswith(b"%PDF")
        path = tmp_path / "r.pdf"
        path.write_bytes(data)
        text = extract_pdf_text(path)
        assert "PROFESSIONAL EXPERIENCE" in text
        assert "20,000+" in text
        lines = [l for l in text.splitlines() if l.strip()]
        one_word = sum(1 for l in lines if len(l.split()) == 1)
        assert one_word / len(lines) < 0.2, "PDF text layer is fragmented"

    def test_build_resume_formats_and_filenames(self):
        pdf, name, mime = build_resume(PROFILE, "pdf", suffix="Acme Corp")
        assert name == "Ada_Lovelace_Resume_Acme_Corp.pdf" and mime == "application/pdf"
        docx, name, _ = build_resume(PROFILE, "docx")
        assert name.endswith(".docx") and docx[:2] == b"PK"
        txt, name, mime = build_resume(PROFILE, "TXT")
        assert b"CORE SKILLS" in txt and mime.startswith("text/plain")
        with pytest.raises(ValueError):
            build_resume(PROFILE, "odt")

    def test_pdf_reflects_tailoring(self, tmp_path):
        data, _, _ = build_resume(PROFILE, "pdf", {"tailored_summary": "UNIQUE-TAILORED-SUMMARY"})
        p = tmp_path / "t.pdf"
        p.write_bytes(data)
        assert "UNIQUE-TAILORED-SUMMARY" in extract_pdf_text(p)

    def test_safe_filename(self):
        r = resume_from_profile(PROFILE)
        assert safe_filename(r, "", "docx") == "Ada_Lovelace_Resume.docx"


# ─── resume_path_for_job ─────────────────────────────────────────────────────

class TestResumePathForJob:
    JOB = {"id": "job123", "company": "Acme & Co", "title": "Eng"}

    def test_falls_back_without_structured_data(self):
        assert resume_path_for_job({"resume_path": "base.pdf"}, self.JOB, {"tailored_summary": "x"}) == "base.pdf"

    def test_falls_back_when_disabled(self):
        p = {**PROFILE, "resume_path": "base.pdf",
             "resume": {**PROFILE["resume"], "use_tailored_when_applying": False}}
        assert resume_path_for_job(p, self.JOB, {"tailored_summary": "x"}) == "base.pdf"

    def test_falls_back_without_tailoring(self):
        p = {**PROFILE, "resume_path": "base.pdf"}
        assert resume_path_for_job(p, self.JOB, {}) == "base.pdf"

    def test_writes_tailored_pdf(self, tmp_path, monkeypatch):
        import utils.usercontext as uc
        monkeypatch.setattr(uc, "resumes_dir", lambda: tmp_path)
        p = {**PROFILE, "resume_path": "base.pdf"}
        out = resume_path_for_job(p, self.JOB, {"tailored_summary": "Tailored for Acme"})
        assert out == str(tmp_path / "tailored" / "job123.pdf")
        assert "Tailored for Acme" in extract_pdf_text(out)


# ─── ATS checker ─────────────────────────────────────────────────────────────

class TestAtsChecker:
    def test_extract_keywords_finds_requirements(self):
        kws = extract_keywords(POSTING)
        low = {k.lower() for k in kws}
        for expected in ("typescript", "react", "node.js", "next.js", "docker", "ci/cd", "postgresql", "websockets"):
            assert expected in low, f"{expected} missing from {kws}"
        assert "requirements" not in low and "experience" not in low

    def test_keywords_present_is_alias_aware(self):
        matched, missing = keywords_present("Built NodeJS services on Google Cloud Platform with web sockets",
                                            ["node.js", "gcp", "websockets", "kubernetes"])
        assert matched == ["node.js", "gcp", "websockets"]
        assert missing == ["kubernetes"]

    def test_check_ats_scores_and_reports_missing(self):
        text = render_text(resume_from_profile(PROFILE))
        report = check_ats(text, POSTING)
        assert 0 <= report["score"] <= 100
        assert report["keyword_score"] is not None
        assert "kubernetes" in {k.lower() for k in report["missing_keywords"]}
        assert "typescript" in {k.lower() for k in report["matched_keywords"]}
        assert report["grade"] in "ABCD"
        assert report["stats"]["sections_found"]  # experience/education/skills detected

    def test_format_checks_flag_fragmented_pdf_text(self):
        fragmented = "\n".join("PRABHJOT SINGH CHADHA EXPERIENCE Software Engineer built things".split())
        report = check_ats(fragmented)
        assert report["keyword_score"] is None
        assert any("one word per line" in i["message"] for i in report["issues"])
        assert any(i["message"].startswith("No email") for i in report["issues"])

    def test_good_resume_has_few_high_issues(self):
        report = check_ats(render_text(resume_from_profile(PROFILE)))
        assert not [i for i in report["issues"] if i["severity"] == "high"]
        assert report["format_score"] >= 70

    def test_check_ats_file_on_generated_pdf(self, tmp_path):
        p = tmp_path / "r.pdf"
        p.write_bytes(render_pdf(resume_from_profile(PROFILE)))
        report = check_ats_file(p, POSTING)
        assert report["source"] == str(p)
        assert report["score"] > 40

    def test_check_ats_file_missing(self, tmp_path):
        report = check_ats_file(tmp_path / "nope.pdf")
        assert report["score"] == 0
        assert report["issues"][0]["severity"] == "high"


# ─── skill pool (customize flow) ─────────────────────────────────────────────

class TestSkillPool:
    def test_pool_merges_resume_profile_and_keywords(self):
        from utils.resume_tailor import skill_pool
        pool = [p.lower() for p in skill_pool(PROFILE)]
        for expected in ("typescript", "python", "gcp", "aws", "docker", "next.js"):
            assert expected in pool
        assert pool.count("typescript") == 1  # deduped across sources

    def test_pool_skills_survive_validation_but_unknown_become_suggestions(self):
        from utils.resume_tailor import _validate_structured, skill_pool
        base = resume_from_profile(PROFILE)
        result = _with_defaults({
            "tailored_skills": [{"category": "Stack", "items": ["Next.js", "TypeScript", "Kubernetes", "Redis"]}],
            "suggested_skills": ["Terraform", "TypeScript"],
        })
        out = _validate_structured(result, base, skill_pool(PROFILE, base))
        # Next.js is in the pool (preferences.keywords) though not in resume.skills;
        # Redis is evidenced in a bullet; Kubernetes is nowhere → suggestion.
        assert out["tailored_skills"] == [{"category": "Stack", "items": ["Next.js", "TypeScript", "Redis"]}]
        assert out["suggested_skills"] == ["Kubernetes", "Terraform"]

    def test_prompt_carries_pool_and_notes(self):
        from utils.resume_tailor import _structured_prompt, skill_pool
        base = resume_from_profile(PROFILE)
        prompt = _structured_prompt(POSTING, base, PROFILE, skill_pool(PROFILE, base), notes="Applying via referral")
        assert "SKILL POOL" in prompt and '"Next.js"' in prompt
        assert "Applying via referral" in prompt
        assert "suggested_skills" in prompt

    def test_extract_keywords_skips_sentence_initial_capitals_and_company(self):
        text = "Ledgerly builds accounting tools. Ship fast with Next.js and Stripe. Fully remote. Ledgerly is hiring."
        kws = {k.lower() for k in extract_keywords(text, exclude=["Ledgerly"])}
        assert "ship" not in kws and "fully" not in kws and "ledgerly" not in kws
        assert "next.js" in kws and "stripe" in kws
